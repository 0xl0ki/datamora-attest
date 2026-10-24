#!/usr/bin/env python3
"""Build the DATAMORA ATTEST demo deck (terminal/ASCII neon style)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── palette (datamora brand: black / green / white, thin) ────────
BG      = RGBColor(0x0B, 0x10, 0x10)
BG2     = RGBColor(0x0A, 0x0F, 0x0F)
GREEN   = RGBColor(0xCA, 0xF9, 0x9C)
TEXT    = RGBColor(0xE4, 0xEA, 0xE5)
MUTED   = RGBColor(0x88, 0x88, 0x88)
DIM     = RGBColor(0x60, 0x60, 0x60)
CYAN    = GREEN
MAGENTA = GREEN
AMBER   = TEXT
RED     = MUTED
FONT    = "SF Mono"

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

def add_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # top accent bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.05))
    bar.fill.solid(); bar.fill.fore_color.rgb = CYAN; bar.line.fill.background()
    return s

def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf

def para(tf, text, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=6, font=FONT):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font
    return p

def pre(slide, x, y, art, size=10, color=CYAN, align=PP_ALIGN.LEFT):
    """Monospace ASCII block, no wrap, preserved spacing."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(13), Inches(4))
    tf = tb.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = art
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
    return tb

def footer(slide, n):
    tf = box(slide, 0.4, 7.05, 8, 0.35)
    para(tf, f"datamora-attest · demo/{n}", size=10, color=DIM, first=True)
    tf2 = box(slide, 12.0, 7.05, 1.0, 0.35)
    para(tf2, f"{n}/06", size=10, color=DIM, align=PP_ALIGN.RIGHT, first=True)

LOGO = (
"██████╗   █████╗  ████████╗  █████╗  ███╗   ███╗  ██████╗  ██████╗   █████╗\n"
"██╔══██╗ ██╔══██╗ ╚══██╔══╝ ██╔══██╗ ████╗ ████║ ██╔═══██╗ ██╔══██╗ ██╔══██╗\n"
"██║  ██║ ███████║    ██║    ███████║ ██╔████╔██║ ██║   ██║ ██████╔╝ ███████║\n"
"██║  ██║ ██╔══██║    ██║    ██╔══██║ ██║╚██╔╝██║ ██║   ██║ ██╔══██╗ ██╔══██║\n"
"██████╔╝ ██║  ██║    ██║    ██║  ██║ ██║ ╚═╝ ██║ ╚██████╔╝ ██║  ██║ ██║  ██║\n"
"╚═════╝  ╚═╝  ╚═╝    ╚═╝    ╚═╝  ╚═╝ ╚═╝     ╚═╝  ╚═════╝  ╚═╝  ╚═╝ ╚═╝  ╚═╝"
)

# ── Slide 1 · cover ──────────────────────────────────────────────
s = add_slide()
pre(s, 0.5, 0.9, LOGO, size=11, color=CYAN, align=PP_ALIGN.CENTER)
tf = box(s, 1.0, 2.6, 11.3, 0.8)
para(tf, "A T T E S T A T I O N   L A Y E R", size=28, color=TEXT, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 3.5, 11.3, 0.9)
para(tf, "AI 已经在链上替人花钱了 —— 凭什么信它？", size=30, color=CYAN, bold=False, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 4.6, 11.3, 0.6)
para(tf, "让自主 Agent 的每一次链上决策：可验证 · 可授权 · 可追责", size=16, color=MUTED, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 6.4, 11.3, 0.5)
para(tf, "ETH Shenzhen 2026 · 方向 01 Agent 与机器经济 / 03 可验证 AI · Robinhood Chain (L2, chainId 4663)",
     size=12, color=DIM, align=PP_ALIGN.CENTER, first=True)
footer(s, 1)

# ── Slide 2 · problem ────────────────────────────────────────────
s = add_slide()
tf = box(s, 0.7, 0.6, 12, 0.8)
para(tf, "> 问题：当 AI Agent 开始自己管钱", size=26, color=CYAN, bold=False, first=True)
problems = [
    ("01", "它到底干了什么？", "不是它嘴上说干了什么 —— 每步操作留一条不可篡改的回执"),
    ("02", "有没有超出授权？", "策略哈希内嵌在回执里，策略一变哈希就变，旧授权自动失效"),
    ("03", "事后能查、能追责吗？", "一批回执压成 Merkle root 锚上链，带时间戳，删不掉、改不了"),
]
y = 1.9
for n, q, a in problems:
    tf = box(s, 1.0, y, 11.4, 1.3)
    para(tf, f"[{n}]  {q}", size=20, color=GREEN, bold=False, first=True, space_after=2)
    para(tf, f"    {a}", size=14, color=MUTED, space_after=0)
    y += 1.55
tf = box(s, 1.0, 6.6, 11.4, 0.6)
para(tf, "答案不是「更聪明的 AI」，而是「不可篡改的账本 + 授权凭证」。", size=15, color=AMBER, first=True)
footer(s, 2)

# ── Slide 3 · 30s demo walkthrough (the one-pager) ───────────────
s = add_slide()
tf = box(s, 0.7, 0.6, 12, 0.8)
para(tf, "> 30 秒走完五个动作", size=26, color=CYAN, bold=False, first=True)
steps = [
    ("①", "存证", "一次 Agent 执行 → 自验证 receipt", "策略哈希 + 输入指纹 + 链上状态，改一个字哈希就崩"),
    ("②", "验真", "任何人离线复算哈希", "篡改立即被 policy_hash_mismatch 检出"),
    ("③", "付费", "x402 按次付费信号", "HTTP 402 → 支付 → 拿到的是可验证信号"),
    ("④", "上链", "批量 receipt → Merkle root", "一笔 tx 锚定，获得链上时间戳"),
    ("⑤", "链上可验", "合约 verifyProof", "证明单条 receipt 就在链上那个 root 里"),
]
y = 1.7
for n, t, d, sub in steps:
    tf = box(s, 0.9, y, 11.6, 0.95)
    para(tf, f"{n}  {t}", size=17, color=CYAN, bold=False, first=True, space_after=1)
    para(tf, f"    {d}    —  {sub}", size=13, color=MUTED, space_after=0)
    y += 1.02
footer(s, 3)

# ── Slide 4 · how to demo / deployment ───────────────────────────
s = add_slide()
tf = box(s, 0.7, 0.6, 12, 0.8)
para(tf, "> 现场怎么演示？要不要上真链？", size=26, color=CYAN, bold=False, first=True)
tf = box(s, 1.0, 1.7, 11.4, 2.4)
para(tf, "主方案：本地链（已部署，零风险）", size=18, color=GREEN, bold=False, first=True, space_after=4)
para(tf, "  本地 EVM · chainId 4663（镜像 Robinhood Chain）· 合约已部署", size=13, color=MUTED, space_after=2)
para(tf, "  ATTEST (LIVE) 就是真发一笔 tx —— 屏幕滚出 txHash + blockNumber，不是 PPT", size=13, color=MUTED, space_after=2)
para(tf, "  完全离线 · 零 gas · 零资金风险 · 不赌会场 WiFi", size=13, color=MUTED, space_after=0)
tf = box(s, 1.0, 4.3, 11.4, 2.2)
para(tf, "加分项：真链锚定（可选）", size=18, color=AMBER, bold=False, first=True, space_after=4)
para(tf, "  赛前把 AttestationAnchor 也部署到公开测试网 / 主网", size=13, color=MUTED, space_after=2)
para(tf, "  那次 tx hash + 浏览器链接贴进 PPT 当旁证", size=13, color=MUTED, space_after=2)
para(tf, "  现场没网也不慌 —— 本地链照常走，真链截图兜底", size=13, color=MUTED, space_after=0)
footer(s, 4)

# ── Slide 5 · tech highlights ────────────────────────────────────
s = add_slide()
tf = box(s, 0.7, 0.6, 12, 0.8)
para(tf, "> 为什么这不是 PPT，是跑通的机制", size=26, color=CYAN, bold=False, first=True)
rows = [
    ("自实现 Keccak-256", "3 个已知向量校验，EVM 原生"),
    ("排序对 Merkle + 补齐 2 的幂", "off-chain JS 与 on-chain Solidity 同构"),
    ("链上 verifyProof / verifyReceipt", "单条 receipt 的成员证明，合约实跑返回 true"),
    ("自验证 receipt（sha256 规范序列化）", "无需信任注册表，任何第三方离线复算"),
    ("x402 付费信号 + HMAC 回执", "按次付费，拿到的是可验证信号"),
    ("跨语言 canonical（Node == Python）", "同一 digest，坐实「无信任第三方可验」"),
]
y = 1.7
for t, d in rows:
    tf = box(s, 1.0, y, 11.4, 0.75)
    para(tf, f"▪ {t}", size=15, color=TEXT, bold=False, first=True, space_after=1)
    para(tf, f"    {d}", size=12.5, color=MUTED, space_after=0)
    y += 0.82
footer(s, 5)

# ── Slide 6 · close ──────────────────────────────────────────────
s = add_slide()
pre(s, 0.5, 1.6, LOGO, size=9, color=DIM, align=PP_ALIGN.CENTER)
tf = box(s, 1.0, 3.2, 11.3, 1.0)
para(tf, "让 AI 花钱这件事：可信、可追责。", size=30, color=CYAN, bold=False, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 4.5, 11.3, 1.0)
para(tf, "存证 → 聚合 → 上链 → 链上可验", size=18, color=GREEN, align=PP_ALIGN.CENTER, first=True)
tf = box(s, 1.0, 5.9, 11.3, 0.5)
para(tf, "GitHub 已开源 · README 含技术栈与运行方式 · 现场 Demo 见 http://127.0.0.1:8787",
     size=12, color=DIM, align=PP_ALIGN.CENTER, first=True)
footer(s, 6)

out = "deck/datamora-attest-demo.pptx"
import os
os.makedirs("deck", exist_ok=True)
prs.save(out)
print("saved", out)
